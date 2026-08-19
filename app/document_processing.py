"""Enterprise document extraction and structure-aware chunking."""

import csv
import io
import os
import re
from pathlib import Path


def _semantic_chunks(text: str, size: int = 850, overlap: int = 120) -> list[str]:
    lines = [line.strip() for line in text.splitlines()]
    sections: list[tuple[str, list[str]]] = []
    heading, paragraphs = "", []
    for line in lines:
        if not line:
            continue
        is_heading = line.startswith("#") or bool(re.match(r"^(?:第[一二三四五六七八九十0-9]+[章节]|\d+(?:\.\d+)*[.、]?)\s+", line))
        if is_heading:
            if paragraphs:
                sections.append((heading, paragraphs))
            heading, paragraphs = line.lstrip("#").strip(), []
        else:
            paragraphs.append(line)
    if paragraphs or heading:
        sections.append((heading, paragraphs))
    chunks = []
    for heading, items in sections:
        current = heading + "\n" if heading else ""
        for paragraph in items:
            candidate = current + paragraph + "\n"
            if len(candidate) > size and current.strip():
                chunks.append(current.strip())
                tail = current[-overlap:] if overlap else ""
                current = (heading + "\n" if heading else "") + tail + paragraph + "\n"
            else:
                current = candidate
        if current.strip():
            chunks.append(current.strip())
    return chunks


def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(raw)).pages)
    if len(text.strip()) >= 80:
        return text
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(raw)) as pdf:
            pieces = []
            for page in pdf.pages:
                pieces.append(page.extract_text() or "")
                for table in page.extract_tables():
                    pieces.extend(" | ".join(cell or "" for cell in row) for row in table)
            text = "\n".join(pieces)
    except Exception:
        pass
    if len(text.strip()) >= 80 or os.getenv("ENABLE_OCR", "0") != "1":
        return text
    try:
        from pdf2image import convert_from_bytes
        import pytesseract

        return "\n".join(pytesseract.image_to_string(page, lang=os.getenv("OCR_LANG", "chi_sim+eng")) for page in convert_from_bytes(raw, dpi=220))
    except Exception as exc:
        raise ValueError(f"PDF 文本提取不足，且 OCR 不可用: {exc}")


def extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(raw)
    if suffix == ".docx":
        from docx import Document

        document = Document(io.BytesIO(raw))
        tables = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        return "\n".join([paragraph.text for paragraph in document.paragraphs] + tables)
    if suffix in {".html", ".htm"}:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw.decode("utf-8", errors="replace"), "html.parser")
        return "\n".join(element.get_text(" ", strip=True) for element in soup.find_all(["h1", "h2", "h3", "p", "li", "tr"]))
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(raw.decode("utf-8-sig", errors="strict")))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    if suffix == ".pptx":
        from pptx import Presentation

        deck = Presentation(io.BytesIO(raw))
        return "\n".join(f"幻灯片 {index + 1}: " + " ".join(shape.text for shape in slide.shapes if hasattr(shape, "text")) for index, slide in enumerate(deck.slides))
    return raw.decode("utf-8", errors="strict")


def process_document(filename: str, raw: bytes) -> list[str]:
    return _semantic_chunks(extract_text(filename, raw))
